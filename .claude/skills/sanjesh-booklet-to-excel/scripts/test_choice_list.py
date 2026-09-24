# -*- coding: utf-8 -*-
"""Prove the choice list works by computing it — on a COPY, never the deliverable.

verify_xlsx.py checks the formulas are wired to the right columns; this script
checks what they compute. It types sample priorities into copies of the
workbook, has LibreOffice recalculate them, and reads the results back:

  scenario A  rows 10, 500, 40, 3000, 7000 get 1, 2, 2.5, 3, 3; a row with an
              empty توضیحات gets 4; one cell gets «۵» typed as Persian-digit
              TEXT. Expect the order 1, 2, 2.5, 3, 3, 4 with every field equal
              to its source row, «تکراری» on both 3s, a blank (not «0») for the
              empty توضیحات, count 6, remaining MAX-6, the text-entry warning,
              and the summary sheet unchanged.
  scenario B  MAX+1 rows get a priority. Expect the list full, the over-limit
              warning, and remaining -1.

  python test_choice_list.py --xlsx book.xlsx [--final final.json] [--timeout 300]

Needs LibreOffice (`soffice`). Exits 0 on PASS, 1 on FAIL, 2 when LibreOffice
is missing — say so to the user rather than claiming the list was tested.
"""
import argparse, collections, io, json, os, shutil, subprocess, sys, tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

DATA = "رشته محل ها"
PRIORITY = "اولویت انتخاب"
ERRORS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A", "Err:")
LIST_SOURCE = {"اولویت واردشده": PRIORITY}

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""


def recalc(paths, timeout):
    """Recalculate and re-save each workbook in place through LibreOffice."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return "LibreOffice (soffice) not found"
    env = dict(os.environ, SAL_USE_VCLPLUGIN="svp")     # headless, no display
    with tempfile.TemporaryDirectory(prefix="lo-profile-") as prof:
        url = Path(prof).as_uri()
        run = lambda args, t: subprocess.run([soffice, "--headless", "--norestore",
                                              f"-env:UserInstallation={url}"] + args,
                                             capture_output=True, text=True, timeout=t, env=env)
        # A core-only install (no libreoffice-calc package) cannot open xlsx,
        # and the macro below then waits forever instead of failing. Probe first.
        from openpyxl import Workbook
        probe = os.path.join(prof, "probe.xlsx")
        pw = Workbook()
        pw.active["A1"] = "=1+1"
        pw.save(probe)
        try:
            run(["--convert-to", "csv", "--outdir", prof, probe], 120)
        except subprocess.TimeoutExpired:
            return "LibreOffice timed out opening a one-cell workbook"
        if not os.path.exists(os.path.join(prof, "probe.csv")):
            return ("LibreOffice cannot open xlsx — the Calc component is missing "
                    "(Debian/Ubuntu: apt-get install libreoffice-calc)")
        mdir = Path(prof) / "user" / "basic" / "Standard"
        if not mdir.exists():
            return "LibreOffice did not create a usable profile"
        (mdir / "Module1.xba").write_text(MACRO, encoding="utf-8")
        for p in paths:
            before = os.stat(p).st_mtime_ns
            try:
                r = run(["vnd.sun.star.script:Standard.Module1.RecalculateAndSave"
                         "?language=Basic&location=application", str(Path(p).absolute())], timeout)
            except subprocess.TimeoutExpired:
                return f"LibreOffice timed out after {timeout}s on {p} (raise --timeout)"
            if r.returncode != 0 or os.stat(p).st_mtime_ns == before:
                return f"LibreOffice did not recalculate {p}: {r.stderr.strip()[:200]}"
    return None


def layout(wb):
    ws = wb[DATA]
    hdr = [c.value for c in ws[1]]
    ls_name = next(n for n in wb.sheetnames if n.startswith("لیست"))
    ls = wb[ls_name]
    hr = next(r for r in range(1, 12) if ls.cell(row=r, column=1).value == "ردیف")
    lhdr = [c.value for c in ls[hr]]
    n = 0
    while ls.cell(row=hr + n + 1, column=1).value == n + 1:
        n += 1
    return ws, hdr, ls_name, hr, lhdr, n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--final", help="final.json — also checks the summary totals")
    ap.add_argument("--timeout", type=int, default=300)
    a = ap.parse_args()

    wb = load_workbook(a.xlsx)
    ws, hdr, ls_name, hr, lhdr, MAX = layout(wb)
    if PRIORITY not in hdr:
        print("no «اولویت انتخاب» column — the workbook was built with --no-choice-list")
        sys.exit(1)
    pc = hdr.index(PRIORITY) + 1
    N = ws.max_row
    code_c = hdr.index("کد رشته محل") + 1
    desc_c = hdr.index("توضیحات") + 1
    if any(ws.cell(row=r, column=pc).value is not None for r in range(2, N + 1)):
        print("FAIL: the workbook under test already has priorities in it")
        sys.exit(1)

    # scenario A — the booklet-size rows are clamped for small test files
    pick = [(min(r, N), p) for r, p in ((10, 1), (500, 2), (40, 2.5), (3000, 3), (7000, 3))]
    used = {r for r, _ in pick}
    empty_desc = next(r for r in range(2, N + 1)
                      if r not in used and not ws.cell(row=r, column=desc_c).value)
    pick.append((empty_desc, 4))
    used.add(empty_desc)
    text_row = next(r for r in range(N, 1, -1) if r not in used)
    expect = sorted(pick, key=lambda rp: (rp[1], rp[0]))      # ROW() breaks the tie

    tmp = tempfile.mkdtemp(prefix="choice-test-")
    fa = os.path.join(tmp, "scenario-a.xlsx")
    fb = os.path.join(tmp, "scenario-b.xlsx")
    for r, p in pick:
        ws.cell(row=r, column=pc, value=p)
    ws.cell(row=text_row, column=pc, value="۵")
    wb.save(fa)
    wb = load_workbook(a.xlsx)
    ws = wb[DATA]
    over = list(range(2, 2 + MAX + 1))
    for i, r in enumerate(over, start=1):
        ws.cell(row=r, column=pc, value=i)
    wb.save(fb)

    print(f"recalculating 2 copies of {os.path.basename(a.xlsx)} with LibreOffice …")
    err = recalc([fa, fb], a.timeout)
    if err:
        print("SKIPPED:", err)
        print("The list was NOT tested — say so to the user.")
        sys.exit(2)

    fails = []
    def check(ok, msg):
        if not ok:
            fails.append(msg)

    # ---- scenario A
    v = load_workbook(fa, data_only=True)
    d, ls = v[DATA], v[ls_name]
    col = {h: j for j, h in enumerate(lhdr, start=1)}
    for s in v.sheetnames:
        for row in v[s].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERRORS):
                    fails.append(f"error value {c.value} at {s}!{c.coordinate}")
    for k, (src_row, prio) in enumerate(expect, start=1):
        r = hr + k
        check(ls.cell(row=r, column=col["اولویت واردشده"]).value == prio,
              f"list row {k}: priority {ls.cell(row=r, column=col['اولویت واردشده']).value!r} != {prio}")
        for h in lhdr:
            src = LIST_SOURCE.get(h, h)
            if src not in hdr:
                continue
            want = d.cell(row=src_row, column=hdr.index(src) + 1).value
            got = ls.cell(row=r, column=col[h]).value
            if want in (None, ""):
                check(got in (None, ""), f"list row {k} {h!r}: {got!r} for an empty source (the «0» bug)")
            else:
                check(str(got) == str(want), f"list row {k} {h!r}: {got!r} != source {want!r}")
        tie = sum(1 for _, p in expect if p == prio) > 1
        w = ls.cell(row=r, column=col["هشدار"]).value or ""
        check((w == "تکراری") == tie, f"list row {k}: warning {w!r}, tie={tie}")
    for k in range(len(expect) + 1, MAX + 1):
        vals = [ls.cell(row=hr + k, column=j).value for j in range(2, len(lhdr) + 1)]
        check(all(x in (None, "") for x in vals), f"list row {k} should be empty: {vals[:3]}")
    count, remain, warn = ls["C3"].value, ls["E3"].value, ls["F3"].value or ""
    check(count == len(pick), f"count {count} != {len(pick)}")
    check(remain == MAX - len(pick), f"remaining {remain} != {MAX - len(pick)}")
    check("عدد نیست" in warn and "1" in warn, f"no text-entry warning: {warn!r}")
    check("بیش از" not in warn, f"over-limit warning with {len(pick)} rows: {warn!r}")
    empty_k = next(k for k, (r, _) in enumerate(expect, start=1) if r == empty_desc)
    got_desc = ls.cell(row=hr + empty_k, column=col["توضیحات"]).value
    print(f"A: order {[p for _, p in expect]} · codes "
          f"{[ls.cell(row=hr + k, column=col['کد رشته محل']).value for k in range(1, len(expect) + 1)]}")
    print(f"A: count={count} remaining={remain} · empty توضیحات shows {got_desc!r} · warning: {warn}")

    if a.final:
        rows = json.load(open(a.final, encoding="utf-8"))
        sm = v["خلاصه آماری"]
        total_cap = sum(r["ظرفیت کل"] for r in rows)
        totals = [(sm.cell(row=r, column=2).value, sm.cell(row=r, column=3).value)
                  for r in range(1, sm.max_row + 1) if sm.cell(row=r, column=1).value == "جمع"]
        check(totals and all(t == (len(rows), total_cap) for t in totals),
              f"summary totals {totals} != {(len(rows), total_cap)}")
        nahve = collections.Counter(r["نحوه پذیرش"] for r in rows)
        for r in range(1, sm.max_row + 1):
            k = sm.cell(row=r, column=1).value
            if k in nahve:
                check(sm.cell(row=r, column=2).value == nahve[k], f"summary {k!r} count off")
        print(f"A: summary totals {totals[0] if totals else None} (final.json: {len(rows)}, {total_cap})")

    # ---- scenario B
    v = load_workbook(fb, data_only=True)
    ls = v[ls_name]
    check(ls["C3"].value == MAX + 1, f"B: count {ls['C3'].value} != {MAX + 1}")
    check(ls["E3"].value == -1, f"B: remaining {ls['E3'].value} != -1")
    check("بیش از" in (ls["F3"].value or ""), f"B: no over-limit warning: {ls['F3'].value!r}")
    check(ls.cell(row=hr + MAX, column=col["اولویت واردشده"]).value == MAX,
          f"B: last list row shows {ls.cell(row=hr + MAX, column=col['اولویت واردشده']).value!r}")
    code_last = v[DATA].cell(row=over[MAX - 1], column=code_c).value
    check(str(ls.cell(row=hr + MAX, column=col["کد رشته محل"]).value) == str(code_last),
          "B: last list row is not the MAX-th priority's code")
    print(f"B: count={ls['C3'].value} remaining={ls['E3'].value} · warning: {ls['F3'].value}")

    shutil.rmtree(tmp, ignore_errors=True)
    for f in fails[:20]:
        print("  FAIL:", f)
    print(f"checks failed: {len(fails)}")
    print("PASS" if not fails else "FAIL")
    sys.exit(0 if not fails else 1)


if __name__ == "__main__":
    main()
