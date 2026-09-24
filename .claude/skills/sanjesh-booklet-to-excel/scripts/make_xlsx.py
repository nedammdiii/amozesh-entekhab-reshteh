# -*- coding: utf-8 -*-
"""Write the deliverable workbook: data sheet + summary + legend.

  python make_xlsx.py --final final.json --out book.xlsx --title "..."

The summary sheet uses COUNTIF/SUMIF so it recalculates when the data changes.
openpyxl writes formulas without cached values, so they read back as None until
Excel opens the file - that is expected, not a bug.
"""
import argparse, sys, io, json, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

_ap = argparse.ArgumentParser()
_ap.add_argument("--final", required=True)
_ap.add_argument("--out", required=True)
_ap.add_argument("--title", default="دفترچه انتخاب رشته - آزمون سراسری")
ARGS = _ap.parse_args()
OUT = ARGS.out

rows = json.load(open(ARGS.final, encoding="utf-8"))

COLS = [
    ("کد رشته محل", 13),
    ("عنوان رشته", 42),
    ("نحوه پذیرش", 20),
    ("دوره تحصیلی", 16),
    ("جنس پذیرش", 12),
    ("استان", 18),
    ("دانشگاه / مؤسسه", 52),
    ("نوع دانشگاه", 30),
    ("زیرمجموعه", 26),
    ("ظرفیت نیمسال اول", 10),
    ("ظرفیت نیمسال دوم", 10),
    ("ظرفیت کل", 9),
    ("ظرفیت زن", 9),
    ("ظرفیت مرد", 9),
    ("شرط بومی / سهمیه", 46),
    ("توضیحات", 60),
    ("بخش دفترچه", 40),
    ("منبع دوره تحصیلی", 18),
    ("صفحه PDF", 9),
]
NUMERIC = {"ظرفیت نیمسال اول", "ظرفیت نیمسال دوم", "ظرفیت کل", "ظرفیت زن", "ظرفیت مرد", "صفحه PDF"}

FONT = "Tahoma"
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, size=10)
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()

# ------------------------------------------------------------------ data ----
ws = wb.active
ws.title = "رشته محل ها"
ws.sheet_view.rightToLeft = True

ws.append([c[0] for c in COLS])
for i, (name, width) in enumerate(COLS, start=1):
    cell = ws.cell(row=1, column=i)
    cell.font = HDR_FONT
    cell.fill = HDR_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = width
ws.row_dimensions[1].height = 34

for r in rows:
    ws.append([r[c[0]] for c in COLS])

last = ws.max_row
for row in ws.iter_rows(min_row=2, max_row=last, max_col=len(COLS)):
    for cell in row:
        cell.font = BODY_FONT
        cell.border = BORDER
        header = COLS[cell.column - 1][0]
        if header in NUMERIC:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif header in ("نحوه پذیرش", "دوره تحصیلی", "جنس پذیرش", "استان", "کد رشته محل"):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=False)

tbl = Table(displayName="Reshtehha", ref=f"A1:{get_column_letter(len(COLS))}{last}")
tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(tbl)
ws.freeze_panes = "A2"

# --------------------------------------------------------------- summary ----
sm = wb.create_sheet("خلاصه آماری")
sm.sheet_view.rightToLeft = True
sm.column_dimensions["A"].width = 46
sm.column_dimensions["B"].width = 16
sm.column_dimensions["C"].width = 18

DATA = "'رشته محل ها'"
COL_LETTER = {name: get_column_letter(i) for i, (name, _) in enumerate(COLS, start=1)}


def block(title, field, values, row):
    c = sm.cell(row=row, column=1, value=title)
    c.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
    c.fill = HDR_FILL
    for col, head in ((2, "تعداد کدرشته"), (3, "مجموع ظرفیت")):
        h = sm.cell(row=row, column=col, value=head)
        h.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
        h.fill = HDR_FILL
        h.alignment = Alignment(horizontal="center")
    letter = COL_LETTER[field]
    cap = COL_LETTER["ظرفیت کل"]
    rng = f"{DATA}!${letter}$2:${letter}${last}"
    caprng = f"{DATA}!${cap}$2:${cap}${last}"
    r = row + 1
    for v in values:
        sm.cell(row=r, column=1, value=v).font = BODY_FONT
        f1 = sm.cell(row=r, column=2, value=f'=COUNTIF({rng},A{r})')
        f2 = sm.cell(row=r, column=3, value=f'=SUMIF({rng},A{r},{caprng})')
        for f in (f1, f2):
            f.font = BODY_FONT
            f.alignment = Alignment(horizontal="center")
        r += 1
    t = sm.cell(row=r, column=1, value="جمع")
    t.font = Font(name=FONT, size=10, bold=True)
    for col in (2, 3):
        s = sm.cell(row=r, column=col,
                    value=f"=SUM({get_column_letter(col)}{row+1}:{get_column_letter(col)}{r-1})")
        s.font = Font(name=FONT, size=10, bold=True)
        s.alignment = Alignment(horizontal="center")
    return r + 2


def order(field):
    c = collections.Counter(r[field] for r in rows)
    return [k for k, _ in c.most_common()]


r = 1
title = sm.cell(row=r, column=1, value="خلاصه آماری " + ARGS.title)
title.font = Font(name=FONT, size=12, bold=True)
r += 2
for label, field in (("نحوه پذیرش", "نحوه پذیرش"),
                     ("نوع دانشگاه", "نوع دانشگاه"),
                     ("دوره تحصیلی", "دوره تحصیلی"),
                     ("جنس پذیرش", "جنس پذیرش"),
                     ("زیرمجموعه (وزارتخانه)", "زیرمجموعه"),
                     ("استان", "استان")):
    r = block(label, field, order(field), r)

# ---------------------------------------------------------------- legend ----
lg = wb.create_sheet("راهنما")
lg.sheet_view.rightToLeft = True
lg.column_dimensions["A"].width = 26
lg.column_dimensions["B"].width = 110

def _uniq(col):
    """Distinct values of a column, most common first — keeps the legend honest."""
    return [k for k, _ in collections.Counter(r[col] for r in rows).most_common() if k]


NOTES = [
    ("منبع", ARGS.title + " — استخراج خودکار از فایل PDF دفترچه "
             f"(صفحات {min(r['صفحه PDF'] for r in rows)} تا {max(r['صفحه PDF'] for r in rows)} فایل PDF)"),
    ("تعداد ردیف", f"{len(rows)} کدرشته‌محل (هر ردیف = یک کدرشته‌محل یکتا)"),
    ("", ""),
    ("کد رشته محل", "کد پنج‌رقمی مندرج در ستون «کدرشته محل» دفترچه."),
    ("عنوان رشته", "دقیقاً مطابق ستون «عنوان رشته» دفترچه."),
    ("نحوه پذیرش", "«با آزمون» یا «صرفا با سوابق تحصیلی» - مطابق ستون دفترچه."),
    ("دوره تحصیلی", "مقادیر موجود در این فایل: " + "، ".join(_uniq("دوره تحصیلی")) + ". "
                    "ستون «منبع دوره تحصیلی» نشان می‌دهد این مقدار مستقیم از ستون دفترچه آمده یا از عنوان بخش دفترچه استنتاج شده است."),
    ("جنس پذیرش", "«زن و مرد» / «فقط زن» / «فقط مرد» - بر اساس دو ستون زن و مرد در دفترچه "
                  "(خط تیره یعنی آن جنسیت پذیرش ندارد)."),
    ("استان", "استان محل تحصیل دانشگاه. برای جدول‌های سهمیه‌ای که در آن‌ها استان بالای جدول نیامده، "
              "استان از روی نام دانشگاه و جدول‌های اصلی همان دانشگاه استخراج شده است."),
    ("دانشگاه / مؤسسه", "نام کامل دانشگاه/دانشکده/مؤسسه به همراه محل تحصیل (در صورت درج در دفترچه)."),
    ("نوع دانشگاه", "مقادیر موجود در این فایل: " + "، ".join(_uniq("نوع دانشگاه")) + "."),
    ("زیرمجموعه", "وزارتخانهٔ متولی: " + "، ".join(_uniq("زیرمجموعه")) + "."),
    ("ظرفیت نیمسال اول/دوم", "عدد ستون «ظرفیت پذیرش نیمسال». خط تیره در دفترچه = صفر."),
    ("ظرفیت کل", "جمع ظرفیت نیمسال اول و دوم."),
    ("ظرفیت زن / ظرفیت مرد", "فقط برای ردیف‌هایی که دفترچه ظرفیت را به تفکیک جنسیت عدد زده است؛ "
                             "در بقیه ردیف‌ها خالی است."),
    ("شرط بومی / سهمیه", "برای جدول‌های سهمیه‌ای (بومی، مناطق محروم، بلایای طبیعی، مصوبه افزایش ظرفیت) "
                         "متن عنوان جدول در دفترچه."),
    ("توضیحات", "ستون «توضیحات» دفترچه (مصاحبه، خوابگاه، تعهد خدمت، بورس و ...)."),
    ("بخش دفترچه", "بخشی از دفترچه که ردیف از آن استخراج شده است."),
    ("صفحه PDF", "شماره صفحه در فایل PDF (نه شماره چاپ‌شده روی صفحه)."),
    ("", ""),
    ("نکته", "استخراج به‌صورت خودکار و بر اساس خطوط جدول و مختصات دقیق هر سلول در PDF انجام شده است؛ "
             f"هر {len(rows)} کد موجود در صفحات جدول‌ها استخراج و با شمارش کدهای هر صفحه "
             "صحت‌سنجی شده است (هیچ کدی جا نیفتاده)."),
]
hdr = lg.cell(row=1, column=1, value="راهنمای ستون‌ها")
hdr.font = Font(name=FONT, size=12, bold=True, color="FFFFFF")
hdr.fill = HDR_FILL
lg.cell(row=1, column=2).fill = HDR_FILL
for i, (k, v) in enumerate(NOTES, start=2):
    a = lg.cell(row=i, column=1, value=k)
    b = lg.cell(row=i, column=2, value=v)
    a.font = Font(name=FONT, size=10, bold=True)
    b.font = BODY_FONT
    a.alignment = Alignment(horizontal="right", vertical="top")
    b.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)

wb.save(OUT)
print("saved:", OUT, "rows:", len(rows))
