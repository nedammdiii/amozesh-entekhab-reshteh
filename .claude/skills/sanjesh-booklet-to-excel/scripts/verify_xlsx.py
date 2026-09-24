# -*- coding: utf-8 -*-
"""Re-read the produced workbook and diff every cell against final.json.

openpyxl writing a cell is not proof the cell holds what you meant. Run this
before handing the file over — it caught a real bug on the first booklet.

It also checks the choice list statically (no Excel needed): the priority
column ships EMPTY, every helper-key formula points at its own row, and every
INDEX formula on the list sheet reads the data column whose header matches its
own. A shifted column letter there shows the wrong field under every heading
without raising a single #REF — test_choice_list.py then proves the values.

  python verify_xlsx.py --final final.json --xlsx book.xlsx
"""
import argparse, io, json, re, sys
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

PRIORITY = "اولویت انتخاب"
KEY = "کلید کمکی"
# list header -> data-sheet header, where the two differ
LIST_SOURCE = {"اولویت واردشده": PRIORITY}
INDEX_RE = re.compile(r"INDEX\('([^']+)'!\$([A-Z]+)\$2:\$([A-Z]+)\$(\d+),")


def check_list(wb, data, hdr, last, say):
    """Static checks of the choice-list sheet -> number of faults."""
    bad = 0
    names = [n for n in wb.sheetnames if n.startswith("لیست") and n != data.title]
    if len(names) != 1:
        say(f"CHOICE LIST: expected one «لیست … انتخاب» sheet, found {names}")
        return 1
    ls = wb[names[0]]
    hr = next((r for r in range(1, 12) if ls.cell(row=r, column=1).value == "ردیف"), None)
    if hr is None:
        say("CHOICE LIST: header row «ردیف» not found")
        return 1
    lhdr = [c.value for c in ls[hr]]
    n = 0
    while ls.cell(row=hr + n + 1, column=1).value == n + 1:
        n += 1
    print(f"choice list: sheet {names[0]!r}, {n} rows, header at row {hr}")
    letter_of = {h: get_column_letter(j) for j, h in enumerate(hdr, start=1)}
    key_col = letter_of.get(KEY)
    ix = lhdr.index("شماره ردیف (کمکی)") + 1 if "شماره ردیف (کمکی)" in lhdr else None
    if ix is None:
        return bad + 1
    for k in range(1, n + 1):
        r = hr + k
        f = ls.cell(row=r, column=ix).value or ""
        want = f"'{data.title}'!${key_col}$2:${key_col}${last}"
        if f.count(want) != 3 or f"SMALL({want},{k})" not in f:
            if bad < 5:
                say(f"  LIST row {k}: helper formula does not read the key column: {f[:90]}")
            bad += 1
        for j, h in enumerate(lhdr, start=1):
            v = ls.cell(row=r, column=j).value
            if not (isinstance(v, str) and "INDEX(" in v):
                continue
            m = INDEX_RE.search(v)
            src = LIST_SOURCE.get(h, h)
            ok = (m and m.group(1) == data.title and m.group(2) == m.group(3)
                  and int(m.group(4)) == last and hdr[_col(m.group(2)) - 1] == src)
            if not ok:
                if bad < 5:
                    say(f"  LIST row {k} col {h!r}: reads the wrong data column: {v[:90]}")
                bad += 1
    return bad


def _col(letters):
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--final", required=True)
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--sheet", default="رشته محل ها")
    a = ap.parse_args()

    rows = json.load(open(a.final, encoding="utf-8"))
    wb = load_workbook(a.xlsx)
    print("sheets:", wb.sheetnames)
    ws = wb[a.sheet]
    print(f"dims: {ws.max_row} x {ws.max_column}")

    hdr = [c.value for c in ws[1]]
    print("headers:", hdr)
    last = len(rows) + 1

    bad = 0
    def say(msg):
        print(msg)

    if ws.max_row != len(rows) + 1:
        print(f"ROW COUNT MISMATCH: sheet {ws.max_row - 1} vs json {len(rows)}")
        bad += 1
    pcol = hdr.index(PRIORITY) + 1 if PRIORITY in hdr else None
    kcol = hdr.index(KEY) + 1 if KEY in hdr else None
    if pcol:
        p = get_column_letter(pcol)
        key_re = re.compile(rf'^=IF\(ISNUMBER\({p}(\d+)\),{p}(\d+)\+ROW\(\)/\d+,""\)$')
    filled = 0
    for i, r in enumerate(rows, start=2):
        for j, h in enumerate(hdr, start=1):
            got = ws.cell(row=i, column=j).value
            if j == pcol:
                # the deliverable ships with no priorities: test ones must never leak
                if got is not None:
                    filled += 1
                continue
            if j == kcol:
                m = key_re.match(got or "")
                if not (m and int(m.group(1)) == i and int(m.group(2)) == i):
                    if bad < 5:
                        print(f"  KEY row {i}: {got!r}")
                    bad += 1
                continue
            want = r.get(h)
            if want == "":
                want = None
            if got != want:
                if bad < 5:
                    print(f"  MISMATCH row {i} col {h!r}: {got!r} != {want!r}")
                bad += 1
    print(f"cell mismatches: {bad}")

    if pcol:
        if filled:
            print(f"PRIORITY COLUMN NOT EMPTY: {filled} cells — rebuild the file, "
                  "never deliver a test copy")
            bad += 1
        if not kcol:
            print("helper column «کلید کمکی» is missing")
            bad += 1
        t = ws.tables.get("Reshtehha")
        want_ref = f"A1:{get_column_letter(len(hdr))}{last}"
        if t is None or t.ref != want_ref:
            print(f"TABLE REF: {t.ref if t else None} != {want_ref}")
            bad += 1
        list_bad = check_list(wb, ws, hdr, last, say)
        print(f"choice-list formula faults: {list_bad}")
        bad += list_bad
    else:
        print("choice list: off (no «اولویت انتخاب» column)")

    formulas = 0
    for name in wb.sheetnames:
        for row in wb[name].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas += 1
    print(f"formula cells: {formulas} (they compute when Excel opens the file)")
    print("PASS" if bad == 0 else "FAIL")
    sys.exit(0 if bad == 0 else 1)


if __name__ == "__main__":
    main()
