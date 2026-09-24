"""ساخت لیست حداکثر ۱۵۰تایی انتخاب رشته از اکسل دفترچه + پروفایل داوطلب (JSON).

فقط بر اساس جواب‌های خود داوطلب کار می‌کند: شانس قبولی تخمین زده نمی‌شود و شهریه
هیچ کدی را حذف نمی‌کند. ترتیب لیست همان ترتیب علاقهٔ داوطلب است و اگر کدها بیشتر از
سقف بود، جاها بین رشته‌های او تقسیم می‌شود تا هر رشته‌ای که نوشته در لیست بیاید.

نمونه:
    python build_list.py --booklet tajrobi-1405.xlsx --profile profile.json --out-dir out --name sara
    python build_list.py --booklet tajrobi.xlsx --booklet zaban.xlsx --profile p.json   # چند گروه

قالب پروفایل: assets/profile-example.json، کلیدها در assets/profile-keys.md و فرم در assets/telegram-form.md
"""
import argparse
import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from booklet import WITH_EXAM, load_booklet, load_json, norm, province_pattern

# سطح‌بندی پیش‌فرض دانشگاه‌ها (فقط برای ترتیب درون یک رشته؛ پروفایل می‌تواند اضافه کند)
UNI_TIER1 = ["علوم پزشکی تهران", "شهید بهشتی", "علوم پزشکی ایران", "علوم پزشکی شیراز",
             "علوم پزشکی اصفهان", "علوم پزشکی مشهد", "علوم پزشکی تبریز", "علوم توانبخشی",
             "صنعتی شریف", "دانشگاه تهران", "صنعتی امیرکبیر", "علم و صنعت", "صنعتی اصفهان",
             "خواجه نصیر", "علامه طباطبایی"]
UNI_TIER3 = ["دانشکده", "پیام نور", "غیرانتفاعی", "ملی مهارت", "گناباد", "دزفول", "تربت حیدریه",
             "شاهرود", "سبزوار", "زابل", "جهرم", "فسا", "آبادان", "خراسان شمالی", "نیشابور",
             "ایرانشهر", "جیرفت", "علوم پزشکی بم", "یاسوج"]

FLAG_COLS = ["تعهد خدمت", "مصاحبه/شرایط خاص", "فقط بومی", "ظرفیت کم", "محرومیت کنکور بعد"]
OUT_COLS = ["اولویت", "کد رشته محل", "عنوان رشته", "دانشگاه / مؤسسه", "استان", "دوره تحصیلی",
            "جنس پذیرش", "شروع", "ظرفیت کل", "ردیف علاقه"] + FLAG_COLS + [
            "شرط بومی / سهمیه", "توضیحات", "صفحه PDF"]


def as_list(v):
    if v is None or v is False:
        return []
    return v if isinstance(v, list) else [v]


def entry_mask(df, e, allowed_courses):
    m = pd.Series(True, index=df.index)
    if e.get("title"):
        m &= df["_title"] == norm(e["title"])
    if e.get("title_contains"):
        m &= df["_title"].str.contains(norm(e["title_contains"]), regex=False)
    courses = as_list(e.get("courses")) or allowed_courses
    if courses:
        m &= df["دوره تحصیلی"].isin(courses)
    unis = [norm(u) for u in as_list(e.get("universities"))]
    if unis:
        m &= df["_uni"].map(lambda u: any(x in u for x in unis))
    ex = [norm(u) for u in as_list(e.get("exclude_universities"))]
    if ex:
        m &= ~df["_uni"].map(lambda u: any(x in u for x in ex))
    provs = as_list(e.get("provinces"))
    if provs:
        m &= df["استان"].isin(provs)
    starts = as_list(e.get("start"))
    if starts:
        m &= df["شروع"].isin(starts)
    return m


def uni_tier(uni, extra):
    if "(محل تحصیل" in uni or "(دانشکده" in uni:
        return 3
    if any(x in uni for x in UNI_TIER1 + extra.get("1", [])):
        return 1
    if any(x in uni for x in UNI_TIER3 + extra.get("3", [])):
        return 3
    return 2


def rank_score(names, value):
    """امتیاز ۰ تا ۳ بر اساس جایگاه در یک فهرست ترتیبی (اولی = ۳)."""
    for i, n in enumerate(names):
        if n and n in value:
            return 3 * (1 - i / max(len(names), 1))
    return None


def allocate(avail, max_codes):
    """تقسیم جاهای لیست بین رشته‌ها: وزن رشتهٔ بالاتر بیشتر، ولی هر رشته حداقل یک جا.

    avail: {ردیف ورودی: تعداد کد موجود} به ترتیب علاقه. جای خالی رشته‌ای که کد کافی ندارد
    به بقیه می‌رسد.
    """
    order = list(avail)
    n = len(order)
    weight = {g: n - i for i, g in enumerate(order)}  # خطی: اولی n، آخری ۱
    alloc = {g: 0 for g in order}
    remaining = max_codes
    active = [g for g in order if avail[g] > 0]
    while remaining > 0 and active:
        total_w = sum(weight[g] for g in active)
        budget = remaining
        for g in active:
            want = max(1, int(budget * weight[g] / total_w))
            take = min(want, avail[g] - alloc[g], remaining)
            alloc[g] += take
            remaining -= take
            if remaining == 0:
                break
        active = [g for g in active if alloc[g] < avail[g]]
    return alloc


def main():
    ap = argparse.ArgumentParser(description="ساخت لیست ۱۵۰تایی انتخاب رشته")
    ap.add_argument("--booklet", action="append", required=True, help="اکسل دفترچه (برای چند گروه تکرار شود)")
    ap.add_argument("--profile", required=True, help="پروفایل داوطلب (JSON)")
    ap.add_argument("--out-dir", default=".")
    ap.add_argument("--name", default="list", help="پیشوند نام فایل‌های خروجی")
    ap.add_argument("--column-map", help='JSON نگاشت نام ستون‌ها، مثل {"نام ستون در فایل": "کد رشته محل"}')
    a = ap.parse_args()

    prof = load_json(a.profile)
    cmap = load_json(a.column_map) if a.column_map else None
    frames = [load_booklet(b, cmap) for b in a.booklet]
    booklet_warns = [f"{Path(b).name}: {w}" for b, f in zip(a.booklet, frames) for w in f.attrs.get("warnings", [])]
    df = pd.concat(frames, ignore_index=True)
    steps = [("کل ردیف‌های دفترچه", len(df))]

    def step(mask, text):
        nonlocal df
        df = df[mask]
        steps.append((text, len(df)))

    step(df["نحوه پذیرش"] == WITH_EXAM, "فقط «با آزمون» (کدهای فرم ۱۵۰تایی)")
    step(~df["کد رشته محل"].duplicated(), "حذف کدهای تکراری بین گروه‌ها")

    allowed = as_list(prof.get("allowed_courses"))  # خالی = همهٔ دوره‌ها (شهریه فیلتر نمی‌شود)
    if not allowed and prof.get("course_order"):  # اگر فقط ترتیب دوره‌ها آمده، همان‌ها مجازند
        allowed = [x for c in as_list(prof["course_order"]) for x in as_list(c)]
    choices = prof.get("choices") or []
    if not choices:
        raise SystemExit("پروفایل «choices» ندارد؛ رشته‌های مورد علاقهٔ داوطلب را به ترتیب وارد کنید.")

    # هر کد به اولین ورودیِ سازگار (به ترتیب علاقه) تعلق می‌گیرد
    df = df.copy()
    df["_choice"] = -1
    entry_hits = []
    for i, e in enumerate(choices):
        m = entry_mask(df, e, allowed) & (df["_choice"] == -1)
        df.loc[m, "_choice"] = i
        entry_hits.append(int(m.sum()))
    step(df["_choice"] >= 0, "در رشته‌های مورد علاقهٔ داوطلب (و دوره‌های مجاز)")
    df["ردیف علاقه"] = df["_choice"] + 1
    df["_will_enroll"] = df["_choice"].map(lambda i: bool(choices[i].get("will_enroll", True)))

    gender = prof.get("gender", "")
    if gender:
        step(df["جنس پذیرش"].isin(["زن و مرد", f"فقط {gender}"]), "جنسیت سازگار")
    if not prof.get("is_behyar"):
        step(~df["ویژه بهیاران"], "حذف «ویژه بهیاران»")
    if not prof.get("accept_interview"):
        step(~df["مصاحبه/شرایط خاص"], "حذف مصاحبه، بورس و گزینش")
    native = prof.get("native_province", "")
    own = df["_quota"].str.contains(province_pattern(native), regex=True) if native else False
    ok_native = ~df["فقط بومی"] | (own & (~df["سهمیه خاص"] | bool(prof.get("special_native_quota"))))
    step(ok_native, "حذف کدهای «مخصوص بومی» که داوطلب شرایطش را ندارد")
    commit = prof.get("accept_commitment", False)
    if commit is not True:
        step(~df["تعهد خدمت"] | df["استان"].isin(as_list(commit)), "تعهد خدمت فقط با رضایت (و در استان‌های مجاز)")
    if prof.get("accept_start_1406") is False:
        step(df["شروع"] != "مهر ۱۴۰۶", "حذف کدهای شروع مهر ۱۴۰۶")
    if prof.get("accept_start_bahman") is False:
        step(df["شروع"] != "بهمن", "حذف کدهای ورودی بهمن")
    if prof.get("excluded_provinces"):
        step(~df["استان"].isin(prof["excluded_provinces"]), "حذف استان‌های ناخواسته")
    ex_unis = [norm(u) for u in as_list(prof.get("excluded_universities"))]
    if ex_unis:
        step(~df["_uni"].map(lambda u: any(x in u for x in ex_unis)), "حذف دانشگاه‌ها یا شهرهای ناخواسته")
    if prof.get("exclude_codes"):
        step(~df["کد رشته محل"].isin([str(c) for c in prof["exclude_codes"]]), "حذف کدهای دستی")
    if prof.get("retake_next_year"):
        step(~df["محرومیت کنکور بعد"] | df["_will_enroll"],
             "حذف کدهای دارای محرومیت که داوطلب در آن‌ها ثبت‌نام نمی‌کند")

    # ترتیب: فقط خواست داوطلب (ردیف علاقه + استان، دانشگاه و شهر)
    extra_tiers = {k: [norm(x) for x in v] for k, v in (prof.get("university_tiers") or {}).items()}
    prov_pref = [str(x) for x in as_list(prof.get("preferred_provinces"))]
    uni_pref = [norm(x) for x in as_list(prof.get("preferred_universities"))]
    home = norm(prof.get("home_city", ""))
    mode = prof.get("priority", "field")

    # اولویت دوره‌ها: course_order به ترتیب دلخواه داوطلب (هر عضو یک دوره یا فهرستی از دوره‌های هم‌رتبه)
    course_order = as_list(prof.get("course_order"))
    crank = {norm(x): i for i, c in enumerate(course_order) for x in as_list(c)}
    df["_course_rank"] = df["دوره تحصیلی"].map(lambda v: crank.get(norm(v), len(course_order)))
    cmode = prof.get("course_priority", "within_field" if course_order else "tiebreak")
    if not course_order:
        course_weight = 0.0
    elif cmode == "within_field" and mode != "field":
        course_weight = 2.0  # در حالت شهر/متعادل، دورهٔ پایین‌تر امتیاز مکان را کم می‌کند
    elif cmode == "tiebreak":
        course_weight = 0.3
    else:
        course_weight = 0.0  # course_first و within_field در حالت رشته با کلید مرتب‌سازی اعمال می‌شوند

    def place(r):
        s = rank_score(prov_pref, r["استان"]) or 0.0
        u = rank_score(uni_pref, r["_uni"])
        s += u if u is not None else {1: 2.0, 2: 1.0, 3: 0.0}[uni_tier(r["_uni"], extra_tiers)]
        if home and home in r["_uni"] and "(" not in r["_uni"]:
            s += 1
        if course_order:
            s -= course_weight * r["_course_rank"]
        else:
            s -= 0.5 * r["شهریه‌ای"]  # بدون ترتیب دوره: در یک دانشگاه و رشته، روزانه کمی جلوتر
        s -= 1.0 * r["تعهد خدمت"] + 0.3 * (r["شروع"] != "مهر")
        return round(s, 2)

    df["_place"] = df.apply(place, axis=1) if len(df) else []
    if mode == "place":
        keys, asc = ["_place", "_choice"], [False, True]
    elif mode == "balanced":
        df["_key"] = df["_choice"] - 0.5 * df["_place"]
        keys, asc = ["_key", "_choice"], [True, True]
    elif course_order and cmode == "within_field":
        keys, asc = ["_choice", "_course_rank", "_place"], [True, True, False]
    else:
        keys, asc = ["_choice", "_place"], [True, False]
    if course_order and cmode == "course_first":
        keys, asc = ["_course_rank"] + keys, [True] + asc
    df = df.sort_values(keys + ["کد رشته محل"], ascending=asc + [True])
    df["_rank"] = range(len(df))

    # سقف اختیاری هر ورودی («max»)
    capped = pd.Series(False, index=df.index)
    for i, e in enumerate(choices):
        if e.get("max") is not None:
            idx = df.index[df["_choice"] == i]
            capped[idx[int(e["max"]):]] = True
    if capped.any():
        step(~capped, "اعمال سقف تعداد کد هر ورودی (max)")

    # انتخاب حداکثر max_codes: جاها بین رشته‌ها تقسیم می‌شود، بهترین‌های هر رشته می‌مانند
    max_codes = int(prof.get("max_codes", 150))
    if len(df) <= max_codes:
        final = df
    else:
        avail = df.groupby("_choice").size().reindex(range(len(choices)), fill_value=0)
        alloc = allocate({i: int(avail[i]) for i in range(len(choices))}, max_codes)
        final = pd.concat([df[df["_choice"] == i].head(k) for i, k in alloc.items() if k])
        final = final.sort_values("_rank")
    final = final.head(max_codes).copy()
    final.insert(0, "اولویت", range(1, len(final) + 1))

    # هشدارها
    warns = list(booklet_warns)
    per_entry = []
    for i, (e, n) in enumerate(zip(choices, entry_hits)):
        in_pool = int((df["_choice"] == i).sum())
        kept = int((final["_choice"] == i).sum())
        label = e.get("title") or e.get("title_contains")
        per_entry.append((i + 1, label, in_pool, kept))
        if n == 0:
            warns.append(f"ورودی {i + 1} («{label}») با هیچ کدی جور نشد؛ نام دقیق را با "
                         f"inspect_booklet.py --search پیدا کنید یا شرط‌هایش را بازتر کنید.")
        elif in_pool == 0:
            warns.append(f"ورودی {i + 1} («{label}») {n} کد داشت ولی همه با فیلترها حذف شدند.")
    if len(final) < max_codes:
        warns.append(f"لیست {len(final)} کد دارد (کمتر از {max_codes})؛ همهٔ کدهای مناسب رشته‌های داوطلب در لیست آمده است.")
    elif len(df) > len(final):
        warns.append(f"{len(df) - len(final)} کد مناسب دیگر به خاطر سقف {max_codes} کنار رفت "
                     "(از هر رشته، کدهای با ترجیح کمتر).")
    counts = {c: int(final[c].sum()) for c in FLAG_COLS}
    info = [
        (counts["محرومیت کنکور بعد"], "کد محرومیت از کنکور سال بعد دارند (روزانه، یا پزشکی/دندان/دارو/دامپزشکی در هر دوره)."),
        (counts["تعهد خدمت"], "کد تعهد خدمت دارند."),
        (counts["مصاحبه/شرایط خاص"], "کد مصاحبه/گزینش/بورس دارند."),
        (counts["فقط بومی"], "کد «مخصوص بومی» هستند؛ شرط بومی بودن داوطلب را تأیید کنید."),
        (counts["ظرفیت کم"], "کد ظرفیت ۱ یا ۲ نفر دارند."),
        (int((final["شروع"] != "مهر").sum()), "کد شروع غیر از مهر دارند (بهمن یا مهر ۱۴۰۶)."),
    ]
    warns += [f"{n} {t}" for n, t in info if n]

    out = Path(a.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    write_xlsx(out / f"{a.name}.xlsx", final, prof, steps, warns, per_entry)
    write_md(out / f"{a.name}.md", final, prof, steps, warns, per_entry, a.name)

    print("\n".join(f"{t}: {n:,}" for t, n in steps))
    print("کدها در لیست نهایی:", len(final))
    print("سهم هر رشته (ورودی، عنوان، کد مناسب، در لیست):")
    for r in per_entry:
        print("  ", r)
    print("\nهشدارها:\n- " + "\n- ".join(warns))
    print("\nخروجی:", out / f"{a.name}.xlsx", out / f"{a.name}.md")


def cell_value(v):
    if type(v).__name__ in ("bool", "bool_"):
        return "✓" if v else ""
    if hasattr(v, "item"):
        v = v.item()
    return "" if v is None or (isinstance(v, float) and pd.isna(v)) else v


def write_xlsx(path, final, prof, steps, warns, per_entry):
    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    head = PatternFill("solid", fgColor="D9D9D9")
    wb = Workbook()

    def sheet(title, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.sheet_view.rightToLeft = True
        return ws

    def put(ws, r, vals, is_head=False):
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = bold if is_head else font
            cell.alignment = Alignment(horizontal="right", vertical="top",
                                       wrap_text=isinstance(v, str) and len(v) > 45)
            if is_head:
                cell.fill = head

    ws = sheet("لیست", first=True)
    put(ws, 1, OUT_COLS, True)
    for r, (_, row) in enumerate(final.iterrows(), 2):
        put(ws, r, [cell_value(row[c]) for c in OUT_COLS])
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(OUT_COLS))}{max(len(final), 1) + 1}"
    for i, w in enumerate([7, 11, 24, 42, 15, 13, 9, 9, 8, 9, 9, 12, 8, 8, 12, 40, 55, 8], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws = sheet("خلاصه")
    last = len(final) + 1
    col = {c: get_column_letter(OUT_COLS.index(c) + 1) for c in OUT_COLS}
    rng = lambda c: f"'لیست'!${col[c]}$2:${col[c]}${last}"  # noqa: E731
    r = 1
    put(ws, r, ["ردیف علاقه", "رشته", "کد مناسب", "در لیست"], True)
    for i, label, in_pool, _kept in per_entry:
        r += 1
        put(ws, r, [i, label, in_pool, f"=COUNTIF({rng('ردیف علاقه')},A{r})"])
    r += 2
    put(ws, r, ["دوره تحصیلی", "تعداد"], True)
    for t in final["دوره تحصیلی"].drop_duplicates():
        r += 1
        put(ws, r, [t, f"=COUNTIF({rng('دوره تحصیلی')},A{r})"])
    r += 2
    put(ws, r, ["پرچم", "تعداد"], True)
    for f in FLAG_COLS:
        r += 1
        put(ws, r, [f, f'=COUNTIF({rng(f)},"✓")'])
    r += 2
    put(ws, r, ["مراحل فیلتر (خروجی build_list.py)", "ردیف باقی‌مانده"], True)
    for t, n in steps:
        r += 1
        put(ws, r, [t, n])
    for c, w in zip("ABCD", [45, 30, 12, 12]):
        ws.column_dimensions[c].width = w

    ws = sheet("ورودی داوطلب")
    put(ws, 1, ["کلید", "مقدار"], True)
    for i, (k, v) in enumerate(prof.items(), 2):
        put(ws, i, [k, json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else str(v)])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 120

    ws = sheet("هشدارها")
    put(ws, 1, ["هشدار / نکته"], True)
    for i, w in enumerate(warns + ["کدها را پیش از ثبت با دفترچهٔ رسمی و اطلاعیه‌های اصلاحی سنجش تطبیق دهید."], 2):
        put(ws, i, [w])
        ws.cell(row=i, column=1).alignment = Alignment(horizontal="right", wrap_text=True)
    ws.column_dimensions["A"].width = 120

    wb.calculation.fullCalcOnLoad = True  # فرمول‌های «خلاصه» هنگام باز شدن در Excel محاسبه می‌شوند
    wb.save(path)


def write_md(path, final, prof, steps, warns, per_entry, name):
    L = []
    w = L.append
    w(f"# لیست انتخاب رشته — {prof.get('student_name', name)}\n")
    w(f"- تعداد کد: **{len(final)}**")
    w(f"- نسخهٔ اکسل با همهٔ ستون‌ها: `{Path(path).with_suffix('.xlsx').name}`\n")
    w("| ردیف علاقه | رشته | کد مناسب | در لیست |\n|---|---|---|---|")
    for i, label, in_pool, kept in per_entry:
        w(f"| {i} | {label} | {in_pool} | {kept} |")
    w("\n## هشدارها\n")
    for x in warns:
        w(f"- {x}")
    w("- کدها را پیش از ثبت با دفترچهٔ رسمی و اطلاعیه‌های اصلاحی سنجش تطبیق دهید.\n")
    w("## لیست\n")
    w("| # | کد | رشته | دانشگاه / محل تحصیل | دوره | شروع | نکته |")
    w("|---|---|---|---|---|---|---|")
    for _, r in final.iterrows():
        notes = [f for f in ["تعهد خدمت", "مصاحبه/شرایط خاص", "فقط بومی", "ظرفیت کم"] if r[f]]
        w(f"| {r['اولویت']} | {r['کد رشته محل']} | {r['عنوان رشته']} | {r['دانشگاه / مؤسسه']} | "
          f"{r['دوره تحصیلی']} | {r['شروع']} | {'، '.join(notes)} |")
    w("\n## کدها برای ورود سریع (به ترتیب)\n")
    w("```")
    codes = final["کد رشته محل"].tolist()
    for i in range(0, len(codes), 10):
        w("  ".join(f"{i + j + 1:>3}) {c}" for j, c in enumerate(codes[i:i + 10])))
    w("```\n")
    w("## مراحل فیلتر\n")
    w("| مرحله | ردیف باقی‌مانده |\n|---|---|")
    for t, n in steps:
        w(f"| {t} | {n:,} |")
    Path(path).write_text("\n".join(L) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
