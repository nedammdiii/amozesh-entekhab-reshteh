"""ساخت لیست ۱۵۰تایی نمونه از اکسل دفترچهٔ تجربی ۱۴۰۴، طبق konkur-field-selection-guide.md
(بخش‌های ۱۱، ۱۲ و ۱۷) برای یک داوطلب فرضی.

اجرا:
    python build_sample_list.py <مسیر اکسل تجربی ۱۴۰۴> [پوشهٔ خروجی]

خروجی: nemune-list-150-tajrobi-1404.xlsx و nemune-list-150-tajrobi-1404.md
"""
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = sys.argv[1]
OUT_DIR = Path(sys.argv[2] if len(sys.argv) > 2 else ".")
NAME = "nemune-list-150-tajrobi-1404"

# ---------------- داوطلب فرضی (پیوست الف) ----------------
GENDER = "زن"
PROVINCE = "فارس"
HOME_CITY = "شیراز"
# وزن ترجیح هر رشته: برای سه رشتهٔ اول «رشته مهم‌تر از شهر»، برای پیراپزشکی‌ها رشته و شهر هم‌وزن‌ترند
FIELDS = {
    "دکتری عمومی پزشکی": 100,
    "دکتری عمومی دندانپزشکی": 90,
    "دکتری عمومی داروسازی": 82,
    "فیزیوتراپی": 65,
    "علوم آزمایشگاهی": 60,
    "بینایی سنجی": 58,
    "کاردرمانی": 54,
    "علوم تغذیه": 53,
    "تکنولوژی پرتوشناسی": 51,
    "شنوایی شناسی": 50,
}
TOP3 = {"دکتری عمومی پزشکی", "دکتری عمومی دندانپزشکی", "دکتری عمومی داروسازی"}
NEAR_PROVINCES = {"فارس", "بوشهر", "کهگیلویه و بویراحمد", "اصفهان", "یزد", "تهران", "هرمزگان"}
PAID = {"شهریه پرداز", "آزاد تمام وقت", "خودگردان آزاد"}
ABILITY = 7.0  # جایگاه فرضی داوطلب روی مقیاس کیفی سختی (حدود رتبهٔ ۶ هزار در سهمیهٔ منطقه ۲)

# سطح‌بندی تقریبی دانشگاه‌ها (بخش ۱۰ راهنما)
TIER1 = ["علوم پزشکی تهران", "شهید بهشتی", "علوم پزشکی ایران", "علوم پزشکی شیراز",
         "علوم پزشکی اصفهان", "علوم پزشکی مشهد", "علوم پزشکی تبریز", "علوم توانبخشی"]
TIER3 = ["گناباد", "دزفول", "تربت حیدریه", "شاهرود", "سبزوار", "زابل", "جهرم", "فسا",
         "آبادان", "خراسان شمالی", "نیشابور", "ایرانشهر", "جیرفت", "علوم پزشکی بم", "یاسوج", "دانشکده"]
# سختی کیفی هر رشته در دانشگاه سطح ۱ (بدون آمار قبولی؛ فقط برای نمایش روش)
BASE_DIFF = {
    "دکتری عمومی پزشکی": 10, "دکتری عمومی دندانپزشکی": 10, "دکتری عمومی داروسازی": 9,
    "فیزیوتراپی": 7.5, "بینایی سنجی": 7, "علوم آزمایشگاهی": 6.5, "شنوایی شناسی": 5.5,
    "کاردرمانی": 5.5, "تکنولوژی پرتوشناسی": 5.5, "علوم تغذیه": 5.5,
}
TARGET = {"بلندپروازانه": 10, "بلندپروازانهٔ ممکن": 25, "هدف": 65, "امن": 35, "خیلی امن": 15}
LABELS = list(TARGET)


def norm(s):
    s = str(s).replace("ي", "ی").replace("ك", "ک").replace("‌", " ")
    s = s.replace("کهکیلویه", "کهگیلویه")
    return re.sub(r"\s+", " ", s).strip()


def tier(uni):
    if "آزاد اسلامی" in uni:
        return 2
    if "(محل تحصیل" in uni or "(دانشکده" in uni:
        return 3  # پردیس اقماری در شهر دیگر
    if any(t in uni for t in TIER1):
        return 1
    if any(t in uni for t in TIER3):
        return 3
    return 2


def in_home_city(uni):
    return HOME_CITY in uni and "(" not in uni


def label(gap):
    if gap > 1.5:
        return "بلندپروازانه"
    if gap > 0.5:
        return "بلندپروازانهٔ ممکن"
    if gap > -0.5:
        return "هدف"
    if gap > -1.5:
        return "امن"
    return "خیلی امن"


# ---------------- گام ۲ و ۳: مخزن و فیلتر سخت (بخش ۱۷.۸) ----------------
df = pd.read_excel(SRC, sheet_name="رشته محل ها", dtype={"کد رشته محل": str})
steps = [("کل ردیف‌های فایل", len(df))]
state = {"df": df}


def apply(mask_fn, text):
    d = state["df"]
    state["df"] = d[mask_fn(d, d["توضیحات"].fillna("").map(norm),
                            d["شرط بومی / سهمیه"].fillna("").map(norm),
                            d["دانشگاه / مؤسسه"].map(norm))]
    steps.append((text, len(state["df"])))


prov = r"\s*".join(map(re.escape, norm(PROVINCE).split(" ")))
apply(lambda d, n, q, u: d["نحوه پذیرش"] == "با آزمون", "فقط «با آزمون» (کدهای فرم ۱۵۰تایی)")
apply(lambda d, n, q, u: d["عنوان رشته"].isin(FIELDS), "فقط رشته‌های مورد علاقه و پشتیبان")
apply(lambda d, n, q, u: d["جنس پذیرش"].isin(["زن و مرد", f"فقط {GENDER}"]), "جنسیت سازگار")
apply(lambda d, n, q, u: ~n.str.contains("ویژه بهیاران"), "حذف «ویژه بهیاران»")
apply(lambda d, n, q, u: ~n.str.contains("مصاحبه|پیوستها|گزینش|بورس"), "حذف مصاحبه، بورس و گزینش")
apply(lambda d, n, q, u: ~q.str.contains("مخصوص")
      | q.str.contains(rf"استان(?:\s*محروم)?\s*{prov}(?![آ-ی])", regex=True),
      "حذف کدهای «مخصوص بومی» استان‌های دیگر")
apply(lambda d, n, q, u: ~(n.str.contains("تعهد خدمت") | d["بخش دفترچه"].str.contains("تعهد خدمت|افزایش ظرفیت"))
      | (d["استان"] == PROVINCE), "تعهد خدمت فقط در استان فارس")
apply(lambda d, n, q, u: ~d["دوره تحصیلی"].isin(PAID) | (d["عنوان رشته"].isin(TOP3) & u.map(in_home_city)),
      "دورهٔ شهریه‌ای فقط برای ۳ رشتهٔ اول و فقط در شیراز")
apply(lambda d, n, q, u: ~d["دوره تحصیلی"].isin(["غیرانتفاعی", "پیام نور", "مجازی"]),
      "حذف غیرانتفاعی، پیام نور و مجازی")
apply(lambda d, n, q, u: u.map(in_home_city) | ~n.str.contains(r"فاقد خوابگاه(?!\s*برادران)", regex=True),
      "حذف کدهای «فاقد خوابگاه» خارج از شیراز")

pool = state["df"].copy()
pool["_uni"] = pool["دانشگاه / مؤسسه"].map(norm)
pool["_notes"] = pool["توضیحات"].fillna("").map(norm)
pool["سطح دانشگاه"] = pool["_uni"].map(tier)
pool["شهر خود"] = pool["_uni"].map(in_home_city)
pool["تعهد خدمت"] = pool["_notes"].str.contains("تعهد خدمت") | pool["بخش دفترچه"].str.contains("تعهد خدمت|افزایش ظرفیت")
pool["شهریه‌ای"] = pool["دوره تحصیلی"].isin(PAID)
pool["شروع"] = ((pool["ظرفیت نیمسال اول"] == 0) & (pool["ظرفیت نیمسال دوم"] > 0)).map({True: "بهمن", False: "مهر"})
pool["خوابگاه نامطمئن"] = ~pool["شهر خود"] & pool["_notes"].str.contains(
    "عدم تعهد در واگذاری خوابگاه|محدودیت در ارائه خوابگاه|خوابگاه خودگردان")
pool["ظرفیت کم"] = pool["ظرفیت کل"] <= 2
pool["محرومیت کنکور بعد"] = (pool["دوره تحصیلی"] == "روزانه") | pool["عنوان رشته"].isin(TOP3)
fars = pool["استان"] == PROVINCE
top3 = pool["عنوان رشته"].isin(TOP3)

# ---------------- گام ۵: امتیاز ترجیح (فقط خواست داوطلب، بدون شانس) ----------------
pref = pool["عنوان رشته"].map(FIELDS).astype(float)
pref += pool["سطح دانشگاه"].map({1: 6, 2: 3, 3: 0})
pref += pool["شهر خود"] * 4 + (fars & ~pool["شهر خود"]) * 2
pref -= pool["شهریه‌ای"] * 8
pref -= pool["تعهد خدمت"] * 5
pref -= (pool["شروع"] == "بهمن") * 1
pref -= pool["خوابگاه نامطمئن"] * 2
pref -= (~pool["استان"].isin(NEAR_PROVINCES)) * top3.map({True: 3, False: 7})
pool["امتیاز ترجیح"] = pref.round(1)

# ---------------- گام ۴: برچسب شانس کیفی (بخش ۱۱.۴، بدون آمار قبولی) ----------------
diff = pool["عنوان رشته"].map(BASE_DIFF).astype(float)
diff -= pool["سطح دانشگاه"].map({1: 0, 2: 1, 3: 2}) * top3.map({True: 0.8, False: 1.0})
diff -= fars * 1.0                    # مزیت بومی در کدهای علوم پزشکی استان خود
diff -= pool["تعهد خدمت"] * 1.0
diff -= pool["شهریه‌ای"] * 2.0
diff -= (pool["شروع"] == "بهمن") * 0.3
diff += pool["_uni"].str.contains("تهران|شهید بهشتی|ایران") * 0.5
pool["سختی نسبی"] = diff.round(1)
pool["برچسب شانس"] = (diff - ABILITY).map(label)
steps.append(("مخزن نهایی پیش از انتخاب ۱۵۰ کد", len(pool)))

# ---------------- گام ۶ و ۷: انتخاب ۱۵۰ کد و مرتب‌سازی بر اساس ترجیح ----------------
chosen = pd.concat([pool[pool["برچسب شانس"] == lab].nlargest(k, "امتیاز ترجیح") for lab, k in TARGET.items()])
for lab in ["هدف", "امن", "خیلی امن", "بلندپروازانهٔ ممکن"]:  # جای خالی اول به سطوح امن‌تر می‌رسد
    if len(chosen) >= 150:
        break
    rest = pool.drop(chosen.index)
    chosen = pd.concat([chosen, rest[rest["برچسب شانس"] == lab].nlargest(150 - len(chosen), "امتیاز ترجیح")])
final = chosen.sort_values(["امتیاز ترجیح", "سختی نسبی"], ascending=[False, False]).head(150).copy()
final.insert(0, "اولویت", range(1, len(final) + 1))

# ---------------- خروجی اکسل ----------------
COLS = ["اولویت", "کد رشته محل", "عنوان رشته", "دانشگاه / مؤسسه", "استان", "دوره تحصیلی", "شروع",
        "ظرفیت کل", "برچسب شانس", "امتیاز ترجیح", "سختی نسبی", "تعهد خدمت", "شهریه‌ای",
        "خوابگاه نامطمئن", "ظرفیت کم", "محرومیت کنکور بعد", "شرط بومی / سهمیه", "توضیحات", "صفحه PDF"]
FLAGS = ["تعهد خدمت", "شهریه‌ای", "خوابگاه نامطمئن", "ظرفیت کم", "محرومیت کنکور بعد"]
FILLS = {"بلندپروازانه": "F4CCCC", "بلندپروازانهٔ ممکن": "FCE5CD", "هدف": "FFF2CC",
         "امن": "D9EAD3", "خیلی امن": "B6D7A8"}
FONT, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")
LIST_SHEET = "لیست ۱۵۰تایی"


def sheet(wb, title, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.sheet_view.rightToLeft = True
    return ws


def put(ws, row, values, bold=False, fill=None):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = BOLD if bold else FONT
        cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)
        if fill:
            cell.fill = fill


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()
ws = sheet(wb, LIST_SHEET, first=True)
put(ws, 1, COLS, bold=True, fill=HEAD_FILL)
for r, (_, row) in enumerate(final.iterrows(), 2):
    vals = []
    for c in COLS:
        v = row[c]
        if c in FLAGS:
            v = "✓" if v else ""
        elif pd.isna(v):
            v = ""
        elif hasattr(v, "item"):
            v = v.item()
        vals.append(v)
    put(ws, r, vals)
    ws.cell(row=r, column=COLS.index("برچسب شانس") + 1).fill = PatternFill("solid", fgColor=FILLS[row["برچسب شانس"]])
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{len(final) + 1}"
widths(ws, [7, 11, 22, 42, 16, 13, 7, 8, 18, 10, 9, 9, 9, 11, 9, 13, 40, 55, 9])

# خلاصه با فرمول (به شیت لیست ارجاع می‌دهد)
ws = sheet(wb, "خلاصه")
ref = f"'{LIST_SHEET}'!"
col = {c: get_column_letter(COLS.index(c) + 1) for c in COLS}
last = len(final) + 1
r = 1
put(ws, r, ["توزیع برچسب شانس", "تعداد", "درصد"], bold=True, fill=HEAD_FILL)
first_row = r + 1
for lab in LABELS:
    r += 1
    put(ws, r, [lab, f"=COUNTIF({ref}${col['برچسب شانس']}$2:${col['برچسب شانس']}${last},A{r})",
                f"=IF($B${first_row + len(LABELS)}=0,0,B{r}/$B${first_row + len(LABELS)})"])
    ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=FILLS[lab])
    ws.cell(row=r, column=3).number_format = "0.0%"
r += 1
put(ws, r, ["جمع", f"=SUM(B{first_row}:B{r - 1})", f"=SUM(C{first_row}:C{r - 1})"], bold=True)
ws.cell(row=r, column=3).number_format = "0.0%"

r += 2
put(ws, r, ["رشته", "تعداد کد"], bold=True, fill=HEAD_FILL)
for f in FIELDS:
    r += 1
    put(ws, r, [f, f"=COUNTIF({ref}${col['عنوان رشته']}$2:${col['عنوان رشته']}${last},A{r})"])

r += 2
put(ws, r, ["دوره تحصیلی", "تعداد کد"], bold=True, fill=HEAD_FILL)
for dore in sorted(final["دوره تحصیلی"].unique()):
    r += 1
    put(ws, r, [dore, f"=COUNTIF({ref}${col['دوره تحصیلی']}$2:${col['دوره تحصیلی']}${last},A{r})"])

r += 2
put(ws, r, ["پرچم (هشدار)", "تعداد کد"], bold=True, fill=HEAD_FILL)
for flag in FLAGS + ["شروع بهمن"]:
    r += 1
    if flag == "شروع بهمن":
        formula = f'=COUNTIF({ref}${col["شروع"]}$2:${col["شروع"]}${last},"بهمن")'
    else:
        formula = f'=COUNTIF({ref}${col[flag]}$2:${col[flag]}${last},"✓")'
    put(ws, r, [flag, formula])

r += 2
put(ws, r, ["مراحل فیلتر (خروجی اسکریپت build_sample_list.py)", "تعداد ردیف باقی‌مانده"], bold=True, fill=HEAD_FILL)
for text, n in steps:
    r += 1
    put(ws, r, [text, n])
widths(ws, [52, 20, 10])

# پروفایل داوطلب فرضی
ws = sheet(wb, "پروفایل داوطلب")
PROFILE = [
    ("وضعیت", "داوطلب فرضی برای نمایش روش؛ اطلاعات واقعی نیست"),
    ("گروه آزمایشی", "علوم تجربی"),
    ("جنسیت", "زن"),
    ("سهمیه", "منطقه ۲"),
    ("استان بومی / شهر محل زندگی", "فارس / شیراز"),
    ("سطح رتبه (فرضی)", "حدود ۶٬۰۰۰ در سهمیهٔ منطقه ۲ (پزشکی روزانه در دانشگاه‌های کوچک برایش مرزی است)"),
    ("دوره‌های مجاز در کارنامه", "همهٔ دوره‌ها"),
    ("رشته‌ها به ترتیب علاقه", "، ".join(FIELDS)),
    ("رشته‌های ناخواسته", "پرستاری، مامایی، اتاق عمل، هوشبری، فوریت‌ها، بهداشت"),
    ("اولویت رشته یا شهر", "برای پزشکی/دندانپزشکی/داروسازی رشته مهم‌تر است؛ برای پیراپزشکی‌ها رشته و شهر هم‌وزن‌اند"),
    ("شهرها", "همه‌جا قابل قبول؛ فارس و استان‌های نزدیک (بوشهر، کهگیلویه، اصفهان، یزد، تهران، هرمزگان) ترجیح دارند"),
    ("خوابگاه", "خارج از شیراز لازم است؛ کدهای «فاقد خوابگاه» بیرون از شیراز حذف شدند"),
    ("شهریه", "فقط برای پزشکی/دندانپزشکی/داروسازی و فقط در شیراز (کنار خانواده)"),
    ("تعهد خدمت", "فقط در استان فارس قبول دارد"),
    ("مصاحبه، بورس، گزینش", "قبول ندارد"),
    ("ورودی بهمن", "قابل قبول، با ترجیح کمی کمتر"),
    ("کنکور مجدد", "خیر؛ می‌خواهد امسال قبول شود (پس کدهای روزانه مجازند)"),
    ("ریسک‌پذیری", "متوسط"),
]
put(ws, 1, ["مورد", "مقدار"], bold=True, fill=HEAD_FILL)
for i, (k, v) in enumerate(PROFILE, 2):
    put(ws, i, [k, v])
widths(ws, [30, 95])

# روش و هشدارها
ws = sheet(wb, "روش و هشدارها")
NOTES = [
    "این فایل «نمونهٔ آموزشی» است و روش راهنمای konkur-field-selection-guide.md را نشان می‌دهد؛ برای انتخاب رشتهٔ واقعی استفاده نشود.",
    "داده‌ها از دفترچهٔ ۱۴۰۴ است؛ کدهای ۱۴۰۵ ممکن است فرق کنند. هیچ کدی از این لیست را بدون تطبیق با دفترچهٔ ۱۴۰۵ وارد نکنید.",
    "اکسل دفترچه آمار رتبهٔ قبولی ندارد؛ «برچسب شانس» از یک مدل کیفی آمده است: سختی رشته، سطح دانشگاه، بومی بودن، تعهد، شهریه و ورودی بهمن. آمار واقعی قبولی جایگزین آن می‌شود (بخش ۱۱ راهنما).",
    "«امتیاز ترجیح» فقط خواست داوطلب را نشان می‌دهد و ترتیب لیست فقط بر اساس آن است، نه شانس (بخش ۵ راهنما).",
    "در مخزن فقط تعداد کمی کد «هدف» بود؛ طبق راهنما جای خالی به کدهای امن‌تر رسید. نتیجه: لیست محافظه‌کارانه است.",
    "قبولی در هر کد روزانه و در پزشکی/دندانپزشکی/داروسازی (در هر دوره) یعنی محرومیت از کنکور سال بعد؛ ستون «محرومیت کنکور بعد».",
    "کدهای «تعهد خدمت» یک و نیم برابر طول تحصیل تعهد دارند؛ کدهای «شهریه‌ای» شهریهٔ بالا دارند؛ پیش از ثبت با داوطلب مرور شود.",
    "ستون «صفحه PDF» برای بررسی هر کد در دفترچهٔ اصلی است.",
]
put(ws, 1, ["نکته"], bold=True, fill=HEAD_FILL)
for i, t in enumerate(NOTES, 2):
    put(ws, i, [t])
    ws.cell(row=i, column=1).alignment = Alignment(horizontal="right", wrap_text=True)
widths(ws, [120])

OUT_DIR.mkdir(parents=True, exist_ok=True)
wb.calculation.fullCalcOnLoad = True  # فرمول‌های شیت «خلاصه» هنگام باز شدن در Excel محاسبه می‌شوند
wb.save(OUT_DIR / f"{NAME}.xlsx")

# ---------------- خروجی markdown (قالب پیوست ب) ----------------
dist = final["برچسب شانس"].value_counts()
lines = []
w = lines.append
w("# نمونهٔ لیست ۱۵۰تایی انتخاب رشته — تجربی (دادهٔ دفترچهٔ ۱۴۰۴)\n")
w("> **نمونهٔ آموزشی است، نه لیست واقعی.** داوطلب فرضی است، کدها مال دفترچهٔ **۱۴۰۴** هستند و برای ۱۴۰۵ معتبر نیستند. "
  "برچسب‌های شانس **کیفی**‌اند، چون اکسل دفترچه آمار رتبهٔ قبولی ندارد. این فایل فقط نشان می‌دهد هوش مصنوعی "
  "با [راهنما](../konkur-field-selection-guide.md) (بخش‌های ۱۱، ۱۲ و ۱۷) چطور از اکسل دفترچه به لیست ۱۵۰تایی می‌رسد.\n")
w(f"نسخهٔ اکسل با همهٔ ستون‌ها و پرچم‌ها: [`{NAME}.xlsx`]({NAME}.xlsx) — اسکریپت سازنده: [`build_sample_list.py`](build_sample_list.py)\n")
w("## ۱. داوطلب فرضی\n")
w("| مورد | مقدار |\n|---|---|")
for k, v in PROFILE[1:]:
    w(f"| {k} | {v} |")
w("\n## ۲. مراحل فیلتر (بخش ۱۷.۸)\n")
w("| مرحله | ردیف باقی‌مانده |\n|---|---|")
for text, n in steps:
    w(f"| {text} | {n:,} |")
w("\n## ۳. روش امتیازدهی\n")
w("- **امتیاز ترجیح** (فقط خواست داوطلب): وزن رشته + سطح دانشگاه (۶/۳/۰) + شیراز (۴) یا بقیهٔ فارس (۲) "
  "− شهریه‌ای (۸) − تعهد خدمت (۵) − خوابگاه نامطمئن (۲) − ورودی بهمن (۱) − دوری از خانه (۳ برای سه رشتهٔ اول، ۷ برای بقیه).")
w("- **برچسب شانس** (کیفی): سختی پایهٔ رشته، منهای اثر سطح دانشگاه، بومی فارس، تعهد، شهریه و ورودی بهمن، "
  "مقایسه‌شده با جایگاه فرضی داوطلب. با آمار واقعی قبولی باید جایش روش نسبت رتبه (بخش ۱۱.۴) بیاید.")
w("- **انتخاب ۱۵۰ کد:** سهم هدف هر سطح = بلندپروازانه ۱۰، ممکن ۲۵، هدف ۶۵، امن ۳۵، خیلی امن ۱۵. "
  "در هر سطح کدهای با ترجیح بیشتر انتخاب شدند و جای خالی اول به سطوح امن‌تر رسید. **ترتیب نهایی فقط بر اساس امتیاز ترجیح است.**\n")
w("## ۴. خلاصه\n")
w("| برچسب شانس | تعداد |\n|---|---|")
for lab in LABELS:
    w(f"| {lab} | {int(dist.get(lab, 0))} |")
w("\n| رشته | تعداد |\n|---|---|")
for f, n in final["عنوان رشته"].value_counts().items():
    w(f"| {f} | {n} |")
w("\n**هشدارها:**\n")
w(f"- **{int(final['محرومیت کنکور بعد'].sum())} کد** محرومیت از کنکور سال بعد دارند (همهٔ کدهای روزانه و ۳ رشتهٔ پزشکی در هر دوره). چون داوطلب قصد کنکور مجدد ندارد، مشکلی نیست؛ در غیر این صورت باید حذف شوند.")
w(f"- **{int(final['تعهد خدمت'].sum())} کد** تعهد خدمت یک و نیم برابر طول تحصیل در استان فارس دارند.")
w(f"- **{int(final['شهریه‌ای'].sum())} کد** شهریه‌ای‌اند (شهریه‌پرداز علوم پزشکی شیراز و دندانپزشکی آزاد شیراز).")
w(f"- **{int((final['شروع'] == 'بهمن').sum())} کد** ورودی بهمن‌اند.")
w(f"- **{int(final['خوابگاه نامطمئن'].sum())} کد** خارج از شیراز خوابگاه قطعی ندارند («عدم تعهد» یا «محدودیت» در ارائهٔ خوابگاه).")
w(f"- **{int(final['ظرفیت کم'].sum())} کد** ظرفیت ۱ یا ۲ نفر دارند و نوسان آخرین رتبه‌شان زیاد است.")
w("- در مخزن فقط تعداد کمی کد «هدف» بود؛ برای همین لیست به سمت کدهای امن سنگین شده است.\n")
w("## ۵. جدول کامل\n")
w("| # | کد | رشته | دانشگاه / محل تحصیل | دوره | شروع | شانس | هشدار |\n|---|---|---|---|---|---|---|---|")
for _, row in final.iterrows():
    flags = [f for f in ["تعهد خدمت", "شهریه‌ای", "خوابگاه نامطمئن", "ظرفیت کم"] if row[f]]
    w(f"| {row['اولویت']} | {row['کد رشته محل']} | {row['عنوان رشته']} | {row['دانشگاه / مؤسسه']} | "
      f"{row['دوره تحصیلی']} | {row['شروع']} | {row['برچسب شانس']} | {'، '.join(flags)} |")
w("\n## ۶. فهرست کدها برای ورود سریع (به ترتیب)\n")
w("```")
codes = final["کد رشته محل"].tolist()
for i in range(0, len(codes), 10):
    w("  ".join(f"{i + j + 1:>3}) {c}" for j, c in enumerate(codes[i:i + 10])))
w("```\n")
w("## ۷. موارد نیازمند تأیید با دفترچه\n")
w("- همهٔ کدها باید با دفترچهٔ **۱۴۰۵** تطبیق داده شوند (این‌ها کدهای ۱۴۰۴ هستند).")
w("- نوع گزینش کدهای عادی (کشوری یا بومی) در اکسل نیست؛ مزیت بومی فارس فرض شده است (بخش ۱۷.۴).")
w("- شرایط خوابگاه، شهریه و متن کامل تعهدها را از ستون توضیحات و صفحهٔ PDF هر کد بخوانید.")
w("- در ۱۴۰۵، ۳۰٪ ظرفیت پزشکی و دندانپزشکی از مهر ۱۴۰۶ شروع می‌شود و کد جدا دارد (بخش ۶.۴ راهنما)؛ در این نمونهٔ ۱۴۰۴ چنین کدی نیست.\n")
w("## ۸. پیشنهادهای موازی\n")
w("- انتخاب رشتهٔ دانشگاه آزاد (سامانهٔ جدا، تا ۱۰۰ کد) برای رشته‌های پیراپزشکی آزاد.")
w("- پذیرش صرفاً با سوابق تحصیلی (ثبت‌نام جدا) به‌عنوان مسیر پشتیبان.")
(OUT_DIR / f"{NAME}.md").write_text("\n".join(lines) + "\n", encoding="utf-8")

print("\n".join(f"{t}: {n}" for t, n in steps))
print(final["برچسب شانس"].value_counts().to_dict())
print("written:", OUT_DIR / f"{NAME}.xlsx", OUT_DIR / f"{NAME}.md")
